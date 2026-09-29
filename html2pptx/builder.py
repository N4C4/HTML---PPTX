# -*- coding: utf-8 -*-
"""Convierte la geometría medida (o una imagen) en diapositivas de PowerPoint."""
from pptx.dml.color import RGBColor
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Pt

PX = 9525
ALIGN = {'left': PP_ALIGN.LEFT, 'center': PP_ALIGN.CENTER,
         'right': PP_ALIGN.RIGHT, 'justify': PP_ALIGN.JUSTIFY}
DASH = {'dashed': MSO_LINE_DASH_STYLE.DASH, 'dotted': MSO_LINE_DASH_STYLE.ROUND_DOT}


def E(v):
    return int(round(v * PX))


def RGB(c):
    return RGBColor(c[0], c[1], c[2])


def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def set_line(shape, side):
    if not side:
        shape.line.fill.background()
        return
    shape.line.color.rgb = RGB(side['c'])
    shape.line.width = Pt(max(0.25, side['w'] * 0.75))
    if side.get('d') in DASH:
        shape.line.dash_style = DASH[side['d']]


def set_fill(shape, fill):
    if fill:
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGB(fill)
    else:
        shape.fill.background()


def set_pie(shape, a1, a2):
    avLst = shape._element.spPr.find(qn('a:prstGeom')).find(qn('a:avLst'))
    for gd in list(avLst):
        avLst.remove(gd)
    for name, deg in (('adj1', a1), ('adj2', a2)):
        gd = avLst.makeelement(qn('a:gd'),
                               {'name': name, 'fmla': 'val %d' % int(round((deg % 360) * 60000))})
        avLst.append(gd)


def add_side(slide, x1, y1, x2, y2, sd):
    cn = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, E(x1), E(y1), E(x2), E(y2))
    cn.line.color.rgb = RGB(sd['c'])
    cn.line.width = Pt(max(0.25, sd['w'] * 0.75))
    if sd.get('d') in DASH:
        cn.line.dash_style = DASH[sd['d']]
    no_shadow(cn)


def add_text(slide, r):
    text = r.get('text')
    if not text or r['w'] < 1 or r['h'] < 1:
        return
    tb = slide.shapes.add_textbox(E(r['x']), E(r['y']), E(r['w']), E(r['h']))
    tf = tb.text_frame
    size_px = (r.get('size') or 11) / 0.75
    ws = r.get('ws') or 'normal'
    single_line = ('\n' not in text) and (r['h'] <= size_px * 1.7 + 2 or ws in ('nowrap', 'pre'))
    tf.word_wrap = not single_line
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    ratio = 0
    if r.get('lh') and r.get('size'):
        ratio = r['lh'] * 0.75 / r['size']
        if not (0.6 <= ratio <= 3):
            ratio = 0
    for i, ln in enumerate(text.split('\n')):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = ALIGN.get(r.get('align', 'left'), PP_ALIGN.LEFT)
        p.space_before = Pt(0)
        p.space_after = Pt(0)
        if ratio:
            p.line_spacing = ratio
        run = p.add_run()
        run.text = ln
        f = run.font
        f.name = r.get('font') or 'Arial'
        f.size = Pt(r.get('size') or 11)
        f.bold = bool(r.get('bold'))
        f.italic = bool(r.get('italic'))
        f.color.rgb = RGBColor.from_string((r.get('color') or '000000').lstrip('#'))
        if r.get('ls'):
            run.font._element.set('spc', str(int(round(r['ls'] * 0.75 * 100))))
    no_shadow(tb)


def add_shape(slide, r):
    x, y, w, h = r['x'], r['y'], r['w'], r['h']
    sides = r.get('sides') or {}
    present = [(k, sides[k]) for k in ('t', 'r', 'b', 'l') if sides.get(k)]
    if not r.get('fill') and not present:
        return
    minR, maxR = r.get('radius', 0), r.get('radiusMax', 0)
    minWH = min(w, h)
    if minWH > 2 and maxR >= minWH / 2 - 0.75:
        st = MSO_SHAPE.OVAL if abs(w - h) <= 2 else MSO_SHAPE.ROUNDED_RECTANGLE
    elif minWH > 2 and minR >= 2:
        st = MSO_SHAPE.ROUNDED_RECTANGLE
    else:
        st = MSO_SHAPE.RECTANGLE
    shp = slide.shapes.add_shape(st, E(x), E(y), E(w), E(h))
    if st == MSO_SHAPE.ROUNDED_RECTANGLE:
        rr = minR if minR >= 2 else maxR
        shp.adjustments[0] = min(0.5, rr / minWH)
    set_fill(shp, r.get('fill'))
    uniform = len(present) == 4 and all(sd == sides['t'] for _, sd in present)
    if uniform:
        set_line(shp, sides['t'])
    else:
        shp.line.fill.background()
    no_shadow(shp)
    if present and not uniform:
        t, b, l, rr_side = sides.get('t'), sides.get('b'), sides.get('l'), sides.get('r')
        if t:
            add_side(slide, x, y + t['w'] / 2, x + w, y + t['w'] / 2, t)
        if b:
            add_side(slide, x, y + h - b['w'] / 2, x + w, y + h - b['w'] / 2, b)
        if l:
            add_side(slide, x + l['w'] / 2, y, x + l['w'] / 2, y + h, l)
        if rr_side:
            add_side(slide, x + w - rr_side['w'] / 2, y, x + w - rr_side['w'] / 2, y + h, rr_side)


def add_donut(slide, r):
    d = r['donut']
    x, y, w, h = r['x'], r['y'], r['w'], r['h']
    for s in d['slices']:
        pie = slide.shapes.add_shape(MSO_SHAPE.PIE, E(x), E(y), E(w), E(h))
        set_pie(pie, (s['a'] * 3.6 - 90) % 360, (s['b'] * 3.6 - 90) % 360)
        set_fill(pie, s['c'])
        pie.line.fill.background()
        no_shadow(pie)
    ix, iy = x + d['il'], y + d['it']
    iw, ih = w - d['il'] - d['ir'], h - d['it'] - d['ib']
    ov = slide.shapes.add_shape(MSO_SHAPE.OVAL, E(ix), E(iy), E(max(iw, 1)), E(max(ih, 1)))
    set_fill(ov, d.get('bg') or [255, 255, 255])
    if d.get('bc') and d.get('bw'):
        ov.line.color.rgb = RGB(d['bc'])
        ov.line.width = Pt(max(0.25, d['bw'] * 0.75))
    else:
        ov.line.fill.background()
    no_shadow(ov)
    if d.get('content'):
        add_text(slide, {'x': ix, 'y': iy, 'w': iw, 'h': ih, 'text': d['content'],
                         'align': 'center', 'size': d.get('fs') or 11, 'bold': bool(d.get('fw')),
                         'font': r.get('font') or 'Arial', 'color': d.get('fc') or '000000',
                         'italic': False, 'lh': 0})


def add_svg(slide, r):
    s = r.get('svg')
    if not s or not s.get('lines'):
        return
    try:
        vb = [float(v) for v in s['vb'].split()]
    except Exception:
        vb = []
    if len(vb) != 4:
        vb = [0, 0, 100, 100]
    minx, miny, vw, vh = vb
    sx = r['w'] / vw if vw else 1
    sy = r['h'] / vh if vh else 1
    for ln in s['lines']:
        cn = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            E(r['x'] + (ln['x1'] - minx) * sx), E(r['y'] + (ln['y1'] - miny) * sy),
            E(r['x'] + (ln['x2'] - minx) * sx), E(r['y'] + (ln['y2'] - miny) * sy))
        cn.line.color.rgb = RGBColor.from_string(ln['c'].lstrip('#'))
        cn.line.width = Pt(max(0.25, ln['w'] * 0.75))


def add_dump_slide(prs, data, blank):
    """Crea una diapositiva a partir de la geometría medida en el navegador."""
    slide = prs.slides.add_slide(blank)
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGB(data['bg'])
    shapes = texts = 0
    for r in data['els']:
        if r.get('donut'):
            add_donut(slide, r)
            shapes += 1
        elif r.get('svg') is not None and r.get('tag') == 'SVG':
            add_svg(slide, r)
        else:
            has_side = bool(r.get('sides')) and any((r['sides'] or {}).values())
            if r.get('fill') or has_side:
                add_shape(slide, r)
                shapes += 1
            if r.get('text'):
                add_text(slide, r)
                texts += 1
    return slide, shapes, texts


def add_image_slide(prs, path, width, height, blank):
    slide = prs.slides.add_slide(blank)
    slide.shapes.add_picture(path, 0, 0, width=E(width), height=E(height))
    return slide

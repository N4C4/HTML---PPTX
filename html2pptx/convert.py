# -*- coding: utf-8 -*-
"""API pública: convert(inputs, output, width, height, on_progress)."""
import os

from pptx import Presentation

from .builder import E, add_dump_slide, add_image_slide
from .browser import BrowserNotFound  # noqa: F401 (re-export)
from .extract import DumpError, measure_html  # noqa: F401 (re-export)

SUPPORTED = ('.html', '.htm', '.png', '.jpg', '.jpeg')
HTML_EXT = ('.html', '.htm')


class UnsupportedFile(ValueError):
    """Tipo de archivo no soportado (solo HTML, PNG o JPG)."""


def _check_file(path):
    if not os.path.exists(path):
        raise FileNotFoundError('No existe: %s' % path)
    if not path.lower().endswith(SUPPORTED):
        raise UnsupportedFile('Tipo no soportado: %s (solo HTML, PNG o JPG)' % path)


def convert(inputs, output=None, width=1440, height=1024, on_progress=None):
    """Convierte una lista de archivos en un .pptx y devuelve la ruta de salida.

    - .html: se renderiza en un navegador headless y se genera cada elemento como
      forma o cuadro de texto independiente (editable en PowerPoint/Canva).
    - .png/.jpg: cada imagen entra como fondo de una diapositiva.
    """
    if isinstance(inputs, str):
        inputs = [inputs]
    inputs = list(inputs)
    if not inputs:
        raise ValueError('Indique al menos un archivo HTML, PNG o JPG')
    if width < 200 or height < 200:
        raise ValueError('El lienzo debe medir al menos 200 x 200 px')
    for p in inputs:
        _check_file(p)
    say = on_progress or (lambda msg: None)

    if not output:
        base = os.path.splitext(os.path.basename(inputs[0]))[0]
        output = os.path.join(os.path.dirname(os.path.abspath(inputs[0])), base + '.pptx')

    prs = Presentation()
    prs.slide_width = E(width)
    prs.slide_height = E(height)
    blank = prs.slide_layouts[6]

    htmls = [p for p in inputs if p.lower().endswith(HTML_EXT)]
    imgs = [p for p in inputs if p.lower().endswith(('.png', '.jpg', '.jpeg'))]
    if htmls:
        browser = None
        from .browser import find_browser
        browser = find_browser()
        say('Navegador: %s' % browser)
    for h in htmls:
        say('Midiendo %s ...' % os.path.basename(h))
        data = measure_html(h, width, height, browser=browser)
        _, shapes, texts = add_dump_slide(prs, data, blank)
        say('  %s -> %d formas, %d textos' % (os.path.basename(h), shapes, texts))
    for im in imgs:
        add_image_slide(prs, im, width, height, blank)
        say('  %s -> imagen a tamaño de lienzo' % os.path.basename(im))

    prs.save(output)
    say('Guardado: %s (%d diapositiva%s)' % (output, len(prs.slides), '' if len(prs.slides) == 1 else 's'))
    return output
